"""프로그램 저작권 등록 신청용 자료 만들기 (한국저작권위원회).

  python tools/copyright_kit.py --author "홍길동" --made 2026-10-25

결과 (docs/copyright/):
  프로그램명세서.docx       프로그램 개요·기능·구성·환경·규모
  소스코드_전체.pdf          직접 작성한 소스 전체 (오픈소스 라이브러리 제외)
  소스코드_발췌.pdf          앞 25쪽 + 뒤 25쪽 (일부 제출이 허용될 때)
"""
from __future__ import annotations

import argparse
import subprocess
from datetime import date
from pathlib import Path

from fpdf import FPDF

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "copyright"
CODE_FONT = "/usr/share/fonts/truetype/nanum/NanumGothicCoding.ttf"
TEXT_FONT = "/usr/share/fonts/truetype/nanum/NanumGothic.ttf"
PATTERNS = ["engine/localbiseo/*.py", "engine/localbiseo/**/*.py", "app/src/*.ts", "app/src/*.tsx", "app/src/**/*.ts", "app/src/**/*.tsx", "app/src/*.css", "app/src-tauri/src/*.rs", "tools/*.py"]
LINES_PER_PAGE = 62
NAME = "로컬비서 (LocalBiseo)"
VERSION = "0.1.0"


def source_files() -> list[Path]:
    out = subprocess.run(["git", "ls-files", *PATTERNS], cwd=ROOT, capture_output=True, text=True, check=True).stdout.split()
    return sorted({ROOT / p for p in out}, key=lambda p: str(p.relative_to(ROOT)))


class CodePdf(FPDF):
    def __init__(self, title: str):
        super().__init__(format="A4")
        self.title_text = title
        self.add_font("code", fname=CODE_FONT)
        self.add_font("text", fname=TEXT_FONT)
        self.set_auto_page_break(False)
        self.current_file = ""

    def header(self):
        self.set_font("text", size=8)
        self.set_text_color(110)
        self.cell(0, 6, f"{NAME} v{VERSION} 소스코드 — {self.current_file}", new_x="LMARGIN", new_y="NEXT")
        self.set_text_color(0)

    def footer(self):
        self.set_y(-12)
        self.set_font("text", size=8)
        self.set_text_color(110)
        self.cell(0, 6, f"{self.page_no()}", align="C")
        self.set_text_color(0)


def build_pages(files: list[Path]) -> list[tuple[str, list[str]]]:
    """(파일 경로, 줄 목록) 단위의 쪽 목록."""
    pages = []
    for f in files:
        rel = str(f.relative_to(ROOT))
        lines = f.read_text(encoding="utf-8").splitlines() or [""]
        numbered = [f"{i + 1:5d}  {line.expandtabs(4)}" for i, line in enumerate(lines)]
        for start in range(0, len(numbered), LINES_PER_PAGE):
            pages.append((rel, numbered[start : start + LINES_PER_PAGE]))
    return pages


def write_pdf(pages: list[tuple[str, list[str]]], path: Path, note: str = "") -> None:
    pdf = CodePdf(NAME)
    pdf.set_margins(12, 10, 12)
    if note:
        pdf.current_file = "안내"
        pdf.add_page()
        pdf.set_font("text", size=11)
        pdf.multi_cell(0, 7, note)
    for rel, lines in pages:
        pdf.current_file = rel
        pdf.add_page()
        pdf.set_font("code", size=7.2)
        for line in lines:
            # 너무 긴 줄은 잘라서 쪽 밖으로 나가지 않게 한다.
            pdf.cell(0, 4.25, line[:150], new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(path))


def spec_docx(files: list[Path], total_lines: int, author: str, made: str, pages: int) -> Path:
    import docx
    from docx.shared import Pt

    by_lang: dict[str, int] = {}
    for f in files:
        lang = {".py": "Python", ".ts": "TypeScript", ".tsx": "TypeScript(React)", ".rs": "Rust", ".css": "CSS"}.get(f.suffix, f.suffix)
        by_lang[lang] = by_lang.get(lang, 0) + len(f.read_text(encoding="utf-8").splitlines())

    d = docx.Document()
    d.styles["Normal"].font.name = "맑은 고딕"
    d.styles["Normal"].font.size = Pt(10.5)
    d.add_heading("프로그램 명세서", level=1)
    rows = [
        ("프로그램 명칭", NAME),
        ("버전", VERSION),
        ("저작자", author),
        ("창작(완성) 연월일", made),
        ("공표 여부", "미공표 (공표 시 공표일과 공표 방법 기재: 예) 와디즈 펀딩 리워드로 배포)"),
        ("프로그램의 종류", "응용 프로그램 (데스크톱 업무 생산성 소프트웨어)"),
        ("사용 언어", ", ".join(f"{k} {v:,}줄" for k, v in sorted(by_lang.items(), key=lambda x: -x[1]))),
        ("규모", f"소스 파일 {len(files)}개, {total_lines:,}줄 (소스코드 출력물 {pages}쪽)"),
        ("운영 환경", "Windows 10/11 64비트, macOS 12 이상 / RAM 8GB 이상 권장 16GB"),
        ("개발 환경", "Python 3.12, FastAPI, Rust(Tauri 2), React 19, TypeScript 5, Vite"),
    ]
    t = d.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in rows:
        cells = t.add_row().cells
        cells[0].text = k
        cells[1].text = v
    d.add_heading("1. 프로그램의 개요", level=2)
    d.add_paragraph(
        "인터넷 연결 없이 사용자의 PC 안에서만 동작하는 업무 보조 프로그램이다. 녹음 파일을 회의록으로 변환하고, "
        "문서·이미지 속 개인정보를 탐지하여 가린 사본을 만들며, 상품 이미지와 영상을 일괄 처리한다. "
        "여기에 할 일·일정 추출, 클립보드 기록, 화면 문자 인식, 영수증·명함 정리, 맞춤법·문체 변환, 파일 이름 일괄 변경, "
        "중복 파일 정리, PDF 도구, 집중 타이머, 통합 검색 기능을 포함한다."
    )
    d.add_heading("2. 주요 기능", level=2)
    for line in [
        "회의록: 음성 인식 결과를 받아 요약하고 할 일(담당자·기한)과 일정을 추출, docx·txt·md·srt·vtt로 저장",
        "개인정보 가리기: 주민등록번호·전화번호·카드번호(Luhn 검증)·계좌번호·주소·이름 등 탐지, HWPX·DOCX 서식 유지 마스킹, PDF 해당 쪽 래스터화 후 가림, 이미지 OCR 위치 가림",
        "한국어 날짜·시간 표현 해석기: '다음 주 화요일 오후 2시', '11월 3일(화) 2시~4시' 등을 일정으로 변환하고 iCalendar(.ics)로 저장",
        "한국어 문체 변환: 한글 음절 자모 분해·조합으로 해요체·합니다체·공문체 어미를 변환",
        "이미지 일괄 처리: 마켓 규격 프리셋, 워터마크, 배경 제거 / 영상 무음 구간 검출·삭제, 자막 생성",
        "HWP 5.0 바이너리 본문 텍스트 추출기, HWPX 본문 XML 마스킹",
        "통합 검색(SQLite FTS5 trigram), 오프라인 라이선스(Ed25519 서명) 검증",
        "데스크톱 셸: 엔진 프로세스 관리, 전역 단축키, 화면 캡처, 클립보드 감시, 트레이 상주",
    ]:
        d.add_paragraph(line, style="List Bullet")
    d.add_heading("3. 프로그램의 구성", level=2)
    for line in [
        "engine/localbiseo/core: 데이터베이스, 작업 큐, 문서 텍스트 추출(HWP·HWPX·DOCX·PDF), OCR 연동, 검색 색인, 라이선스",
        "engine/localbiseo/modules: 기능별 모듈 16종과 API",
        "app/src: 화면(React) 18개와 공용 구성요소",
        "app/src-tauri/src: 데스크톱 셸(Rust)",
        "tools: 라이선스 발급, 배포용 묶음, 오픈소스 고지문 생성",
    ]:
        d.add_paragraph(line, style="List Bullet")
    d.add_heading("4. 제외 범위 (타인의 저작물)", level=2)
    d.add_paragraph(
        "본 등록 대상은 위 소스 파일에 한한다. 실행에 사용되는 오픈소스 라이브러리(FastAPI, Pillow, pypdf, pypdfium2, "
        "python-docx, faster-whisper, Tesseract, Tauri, React 등)와 AI 모델, 글꼴은 각 권리자의 저작물이며 등록 대상에서 제외한다. "
        "목록과 라이선스는 설치 파일의 THIRD_PARTY_NOTICES.txt에 있다."
    )
    d.add_heading("5. 생성형 AI 활용 여부", level=2)
    d.add_paragraph(
        "본 프로그램의 소스코드 상당 부분은 저작자의 기획·지시·검토 아래 생성형 AI 코딩 도구(Anthropic Claude)를 사용해 작성되었다. "
        "[저작자가 직접 작성하거나 창작적으로 수정한 부분, 기능 기획·구성·선택 등 인간의 창작적 기여를 구체적으로 적는다.]"
    )
    path = OUT / "프로그램명세서.docx"
    d.save(path)
    return path


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--author", default="[저작자 이름]")
    ap.add_argument("--made", default=date.today().isoformat())
    args = ap.parse_args()
    OUT.mkdir(parents=True, exist_ok=True)
    files = source_files()
    total = sum(len(f.read_text(encoding="utf-8").splitlines()) for f in files)
    pages = build_pages(files)
    note = (
        f"{NAME} v{VERSION} 소스코드 출력물\n\n"
        f"- 대상: 직접 작성한 소스 파일 {len(files)}개, {total:,}줄\n"
        "- 제외: 오픈소스 라이브러리, AI 모델, 글꼴, 자동 생성 파일, 테스트 코드\n"
        "- 각 쪽 위에 파일 경로, 각 줄 앞에 줄 번호를 표시했다."
    )
    write_pdf(pages, OUT / "소스코드_전체.pdf", note)
    excerpt = pages[:25] + pages[-25:] if len(pages) > 50 else pages
    write_pdf(excerpt, OUT / "소스코드_발췌.pdf", note + "\n- 발췌본: 앞 25쪽과 뒤 25쪽")
    spec = spec_docx(files, total, args.author, args.made, len(pages))
    print(f"파일 {len(files)}개, {total:,}줄, {len(pages)}쪽 → {OUT}\n{spec}")


if __name__ == "__main__":
    main()
