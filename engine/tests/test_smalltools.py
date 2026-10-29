from __future__ import annotations

import pytest
from PIL import Image

from localbiseo.modules.smalltools import en_to_ko, hangul_amount, ko_to_en


def test_count(client):
    r = client.post("/tools/count", json={"text": "안녕하세요 반가워요.\nabc\n\n둘째 문단"}).json()
    assert r["chars"] == 22 and r["chars_no_space"] == 17 and r["words"] == 5
    assert r["lines"] == 4 and r["paragraphs"] == 2
    assert r["bytes_ko"] == 13 * 2 + 9 and r["bytes_utf8"] == 13 * 3 + 9


@pytest.mark.parametrize("n,formal,reading", [
    (0, "영", "영"),
    (10000, "일만", "만"),
    (1234567, "일백이십삼만사천오백육십칠", "백이십삼만 사천오백육십칠"),
    (100_000_000, "일억", "일억"),
    (3_000_015_000, "삼십억일만오천", "삼십억 만 오천"),
])
def test_hangul_amount(n, formal, reading):
    assert hangul_amount(n, formal=True) == formal
    assert hangul_amount(n) == reading


def test_money_and_vat(client):
    r = client.post("/tools/money", json={"amount": "1,234,567원", "vat": "add"}).json()
    assert r["formal"] == "일금 일백이십삼만사천오백육십칠원정" and r["short"] == "123만 4,567원"
    assert r["vat"] == {"supply": 1234567, "tax": 123456, "total": 1358023}
    r = client.post("/tools/money", json={"amount": "110000", "vat": "split"}).json()
    assert r["vat"] == {"supply": 100000, "tax": 10000, "total": 110000}
    r = client.post("/tools/money", json={"amount": "55555", "vat": "split"}).json()
    assert r["vat"]["supply"] + r["vat"]["tax"] == 55555 and r["vat"]["supply"] == 50505
    assert client.post("/tools/money", json={"amount": "삼만원"}).status_code == 400


def test_date_between_add_dday(client):
    # 2026년 추석 연휴(9/24~26)와 개천절 대체공휴일(10/5)이 낀 구간
    r = client.post("/tools/date", json={"mode": "between", "start": "2026-09-21", "end": "2026-10-09"}).json()
    assert r["days"] == 18 and r["days_inclusive"] == 19
    assert r["workdays"] == 15 - 2 - 1 - 1  # 평일 15일 - 추석 2일(목·금) - 10/5 - 한글날
    assert {h["holiday"] for h in r["holidays"]} >= {"추석", "대체공휴일(개천절)", "한글날"}
    r = client.post("/tools/date", json={"mode": "add", "start": "2026-09-23", "days": 3, "workdays": True}).json()
    assert r["result"] == {"date": "2026-09-30", "weekday": "수", "holiday": None}
    r = client.post("/tools/date", json={"mode": "add", "start": "2026-01-31", "days": 30}).json()
    assert r["result"]["date"] == "2026-03-02" and r["result"]["holiday"] == "대체공휴일(삼일절)"
    r = client.post("/tools/date", json={"mode": "dday", "end": "2026-11-11", "today": "2026-10-29"}).json()
    assert r["label"] == "D-13" and r["days"] == 13
    r = client.post("/tools/date", json={"mode": "dday", "end": "2030-01-01", "today": "2026-10-29"}).json()
    assert r["notes"]
    assert client.post("/tools/date", json={"mode": "between", "start": "어제", "end": "2026-01-01"}).status_code == 400
    names = {h["name"] for h in client.get("/tools/holidays", params={"year": 2026}).json()}
    assert {"노동절", "제헌절", "대체공휴일(광복절)"} <= names


def test_qr(client, tmp_path):
    r = client.post("/tools/qr", json={"text": "https://www.wadiz.kr", "out_path": str(tmp_path / "q.png")}).json()
    img = Image.open(r["dst"])
    assert img.size[0] == img.size[1] and img.size[0] > 200 and r["data_url"].startswith("data:image/png;base64,")
    r = client.post("/tools/qr", json={"kind": "wifi", "ssid": "가게;손님", "password": "pw:123", "out_path": str(tmp_path / "w.png")}).json()
    assert r["content"] == r"WIFI:T:WPA;S:가게\;손님;P:pw\:123;;"
    assert client.post("/tools/qr", json={"text": " "}).status_code == 400


def test_password(client):
    r = client.post("/tools/password", json={"length": 12, "count": 3}).json()
    assert len(r["passwords"]) == 3 and all(len(p) == 12 for p in r["passwords"])
    for p in r["passwords"]:
        assert any(c.isupper() for c in p) and any(c.isdigit() for c in p) and not set(p) & set("Il1O0o")
    r = client.post("/tools/password", json={"length": 8, "upper": False, "digits": True, "symbols": False}).json()
    assert all(c.islower() or c.isdigit() for c in r["passwords"][0]) and r["strength"] == "약함"


@pytest.mark.parametrize("en,ko", [
    ("dkssudgktpdy", "안녕하세요"),
    ("rhkswjd", "관정"),
    ("dlTek", "있다"),
    ("ekfrdmf", "닭을"),
    ("Rkcl", "까치"),
    ("qnpfrwhr", "뷁족"),
    ("Gkqn 123!", "하부 123!"),
    ("ghkdmlwkd", "화의장"),
])
def test_keyboard_roundtrip(en, ko):
    assert en_to_ko(en) == ko
    assert ko_to_en(ko) == en.replace("G", "g")


def test_keyboard_api(client):
    assert client.post("/tools/keyboard", json={"text": "gksrmf"}).json() == {"direction": "en2ko", "text": "한글"}
    assert client.post("/tools/keyboard", json={"text": "ㅗ디ㅣㅐ"}).json()["text"] == "hello"


def test_clean(client):
    text = "• 첫째  줄 \n• 둘째 줄\n\n• 첫째  줄 \n１２３　ＡＢＣ 😀"
    r = client.post("/tools/clean", json={"text": text, "ops": ["bullets", "trim", "spaces", "blank", "dedupe", "halfwidth", "emoji", "number"]}).json()
    assert r["text"] == "1. 첫째 줄\n2. 둘째 줄\n3. 123 ABC"
    pdf = "회의는 다음 주\n화요일에 합니다.\n\n둘째 문단\n이어짐"
    assert client.post("/tools/clean", json={"text": pdf, "ops": ["join_lines"]}).json()["text"] == "회의는 다음 주 화요일에 합니다.\n\n둘째 문단 이어짐"
    assert client.post("/tools/clean", json={"text": "x", "ops": ["없음"]}).status_code == 400


def test_units(client):
    r = {x["unit"]: x["value"] for x in client.post("/tools/units", json={"value": 84, "unit": "㎡"}).json()["results"]}
    assert r["평"] == pytest.approx(25.41, abs=0.01)
    r = {x["unit"]: x["value"] for x in client.post("/tools/units", json={"value": 1, "unit": "근"}).json()["results"]}
    assert r["g"] == 600 and r["돈"] == 160
    r = {x["unit"]: x["value"] for x in client.post("/tools/units", json={"value": 100, "unit": "°F"}).json()["results"]}
    assert r["°C"] == pytest.approx(37.7778, abs=1e-3)
    assert client.post("/tools/units", json={"value": 1, "unit": "광년"}).status_code == 400


def test_check_numbers(client):
    r = client.post("/tools/check", json={"value": "124-81-00998"}).json()  # 삼성전자 사업자번호
    assert r["kind"] == "사업자등록번호" and r["valid"] and r["detail"] == "영리법인 본점"
    assert not client.post("/tools/check", json={"value": "124-81-00999"}).json()["valid"]
    r = client.post("/tools/check", json={"value": "4111 1111 1111 1111"}).json()
    assert r["kind"] == "카드번호" and r["valid"] and r["detail"] == "비자"
    r = client.post("/tools/check", json={"value": "900101-1234567"}).json()
    assert r["kind"] == "주민등록번호" and r["formatted"] == "900101-1234567"
    r = client.post("/tools/check", json={"value": "01012345678"}).json()
    assert r["kind"] == "전화번호" and r["formatted"] == "010-1234-5678"
    assert client.post("/tools/check", json={"value": "abc"}).json()["kind"] == ""


def test_merge_tables(client, tmp_path):
    from openpyxl import Workbook, load_workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["이름", "금액"])
    ws.append(["가게A", 1000])
    wb.save(tmp_path / "1월.xlsx")
    (tmp_path / "2월.csv").write_bytes("이름,메모,금액\n가게B,반품,2000\n\n".encode("cp949"))
    r = client.post("/tools/merge-tables", json={"paths": [str(tmp_path)], "out_path": str(tmp_path / "out" / "합침.xlsx")}).json()
    assert r["rows"] == 2 and [f["rows"] for f in r["files"]] == [1, 1]
    rows = [list(x) for x in load_workbook(r["dst"]).active.iter_rows(values_only=True)]
    assert rows == [["출처", "이름", "금액", "메모"], ["1월.xlsx", "가게A", 1000, None], ["2월.csv", "가게B", "2000", "반품"]]
    assert client.post("/tools/merge-tables", json={"paths": [str(tmp_path / "없음")]}).status_code == 400
