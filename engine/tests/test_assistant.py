from __future__ import annotations

import struct
from datetime import date, datetime, timedelta

from PIL import Image

from localbiseo.core.extract import _hwp_para_text, _hwp_records
from localbiseo.modules.papers import parse_card, parse_receipt, to_vcard
from localbiseo.modules.writing import change_tone, email, spellcheck


# --- 글다듬기 ---
def test_spellcheck_common_mistakes():
    r = spellcheck("몇일 뒤에 할께요. 그건 안되요. 어의없네  정말 . 할수있어요 할때는 금새 끝나요. 오랫만이에요 됬다 할려고 했슴.")
    assert r["text"] == "며칠 뒤에 할게요. 그건 안 돼요. 어이없네 정말. 할 수 있어요 할 때는 금세 끝나요. 오랜만이에요 됐다 하려고 했음."
    assert len(r["changes"]) >= 10


def test_spellcheck_leaves_correct_text():
    text = "왠지 오늘은 할수록 좋아요. 선생님께 드릴 게 있어요. 깨끗이 정리했습니다."
    assert spellcheck(text)["text"] == text


def test_tone_polite_to_formal_and_official():
    t = "회의는 3시에 해요. 자료는 김대리가 보내요. 확인해 주세요. 시간 있나요? 정말 어려워요. 제가 할게요. 그건 아니에요."
    assert change_tone(t, "formal")["text"] == (
        "회의는 3시에 합니다. 자료는 김대리가 보냅니다. 확인해 주십시오. 시간 있습니까? 정말 어렵습니다. 제가 하겠습니다. 그건 아닙니다."
    )
    assert change_tone(t, "official")["text"] == (
        "회의는 3시에 함 자료는 김대리가 보냄 확인해 주시기 바람 시간 있습니까? 정말 어려움 제가 하겠음 그건 아님"
    )


def test_tone_formal_to_polite():
    t = "회의는 3시에 합니다. 자료를 보냅니다. 문제가 있습니다. 좋습니다. 확인하십시오. 사과입니다. 책입니다. 어렵습니다. 했습니다. 드립니다."
    assert change_tone(t, "polite")["text"] == "회의는 3시에 해요. 자료를 보내요. 문제가 있어요. 좋아요. 확인하세요. 사과예요. 책이에요. 어려워요. 했어요. 드려요."


def test_tone_keeps_nuanced_endings():
    t = "좋네요. 그렇죠? 갈까요?"
    assert change_tone(t, "formal")["text"] == t


def test_email_template():
    out = email("자료 보내드려요.", "김민수", "이영희", "로컬컴퍼니")["text"]
    assert out.startswith("안녕하세요, 김민수님.\n로컬컴퍼니 이영희입니다.")
    assert "자료 보내드립니다." in out and out.endswith("로컬컴퍼니 이영희 드림")


# --- 영수증·명함 ---
RECEIPT = """스타벅스 강남역점
사업자번호 220-81-62517
2026-10-20 12:31
아메리카노 2 9,000
카페라떼 1 5,000
합계금액 14,000원
카드승인 14,000
"""


def test_parse_receipt():
    r = parse_receipt(RECEIPT)
    assert r == {"merchant": "스타벅스 강남역점", "date": "2026-10-20", "amount": 14000, "biz_no": "220-81-62517", "category": "식비"}


CARD = """주식회사 로컬컴퍼니
홍길동
마케팅팀 / 팀장
M. 010-1234-5678
T. 02-555-1234  F. 02-555-1235
gildong@localco.kr
www.localco.kr
서울특별시 강남구 테헤란로 123, 5층
"""


def test_parse_card_and_vcard():
    c = parse_card(CARD)
    assert c["company"] == "주식회사 로컬컴퍼니"
    assert c["name"] == "홍길동"
    assert c["title"] == "팀장"
    assert c["mobile"] == "010-1234-5678"
    assert c["phone"] == "02-555-1234" and c["fax"] == "02-555-1235"
    assert c["email"] == "gildong@localco.kr" and c["website"] == "www.localco.kr"
    assert c["address"].startswith("서울특별시 강남구")
    v = to_vcard(c)
    assert "FN:홍길동" in v and "TEL;TYPE=CELL:010-1234-5678" in v and v.endswith("END:VCARD\r\n")


def test_expenses_and_contacts_api(client, tmp_path):
    fields = client.post("/receipts/parse", json={"text": RECEIPT}).json()["fields"]
    client.post("/expenses", json=fields)
    client.post("/expenses", json={"date": "2026-10-21", "merchant": "택시", "amount": 8800, "category": "교통"})
    r = client.post("/expenses/export", json={"month": "2026-10", "out_path": str(tmp_path / "e.xlsx")}).json()
    assert r["count"] == 2 and r["total"] == 22800
    from openpyxl import load_workbook

    ws = load_workbook(tmp_path / "e.xlsx").active
    assert ws["A1"].value == "날짜" and ws["C4"].value == "=SUM(C2:C3)"
    card = client.post("/cards/parse", json={"text": CARD}).json()["fields"]
    client.post("/contacts", json=card)
    r = client.post("/contacts/export", json={"out_path": str(tmp_path / "c.vcf")}).json()
    assert r["count"] == 1
    assert client.get("/search", params={"q": "로컬컴퍼니"}).json()["results"][0]["kind"] == "contact"


# --- 메모·할 일·검색·클립보드 ---
def test_notes_todos_search(client):
    n = client.post("/notes", json={"title": "신제품 아이디어", "body": "와디즈 펀딩 리워드 구성 고민", "tags": "기획"}).json()
    client.put(f"/notes/{n['id']}", json={"title": "신제품 아이디어", "body": "와디즈 펀딩 리워드 구성 확정", "tags": "기획"})
    t = client.post("/todos", json={"text": "펀딩 스토리 초안 쓰기", "due": "2026-10-30"}).json()
    client.patch(f"/todos/{t['id']}", json={"done": True})
    assert client.get("/todos", params={"status": "open"}).json() == []
    assert client.get("/todos", params={"status": "done"}).json()[0]["done_at"]
    res = client.get("/search", params={"q": "리워드 구성"}).json()["results"]
    assert res[0]["kind"] == "note" and "확정" in res[0]["snippet"]
    # 2글자 검색어는 LIKE로 찾는다
    assert any(r["kind"] == "todo" for r in client.get("/search", params={"q": "초안"}).json()["results"])
    client.delete(f"/notes/{n['id']}")
    assert client.get("/search", params={"q": "리워드 구성"}).json()["results"] == []


def test_clipboard_masks_sensitive(client):
    r = client.post("/clipboard", json={"text": "제 번호는 010-1234-5678 입니다"}).json()
    assert r["sensitive"] and r["kinds"] == ["phone"]
    assert client.post("/clipboard", json={"text": "제 번호는 010-1234-5678 입니다"}).json().get("duplicate")
    client.post("/clipboard", json={"text": "그냥 문장"})
    items = client.get("/clipboard").json()
    assert items[0]["text"] == "그냥 문장" and items[1]["text"] == "제 번호는 010-****-5678 입니다"
    client.put("/clipboard/policy", json={"policy": "skip"})
    assert client.post("/clipboard", json={"text": "카드 4111 1111 1111 1111"}).json()["skipped"]
    client.post(f"/clipboard/{items[1]['id']}/pin")
    client.delete("/clipboard")
    assert [c["id"] for c in client.get("/clipboard").json()] == [items[1]["id"]]


def test_focus_report(client, app):
    s = client.post("/focus/start", json={"label": "기획서", "planned_minutes": 25}).json()
    # 시작 시각을 30분 전으로 당겨서 기록
    app.state.ctx.db.execute("UPDATE focus_sessions SET started_at=? WHERE id=?", ((datetime.now() - timedelta(minutes=30)).isoformat(timespec="seconds"), s["id"]))
    client.post(f"/focus/{s['id']}/stop", json={"completed": True})
    today = client.get("/focus/today").json()
    assert today["minutes"] == 30 and today["completed"] == 1
    rep = client.get("/focus/report", params={"week": date.today().isoformat()}).json()
    assert rep["total_minutes"] == 30 and rep["labels"][0]["label"] == "기획서"
    assert "30분" in rep["markdown"] or "0시간 30분" in rep["markdown"]


# --- 회의록 (붙여넣기 → 요약·할 일·내보내기) ---
MEETING = (
    "오늘 회의 안건은 와디즈 펀딩 준비입니다. 목표 금액은 300만 원으로 정했습니다. "
    "김민수 대리가 금요일까지 상세페이지 시안을 공유하기로 했습니다. "
    "리워드는 슈퍼얼리버드 29,000원과 얼리버드 39,000원 두 가지로 확정했습니다. "
    "다음 주 화요일 오후 2시에 촬영 일정을 잡았습니다. "
    "예산 문제는 다음 회의에서 다시 논의하겠습니다. 고객 문의 번호는 010-1234-5678입니다."
)


def test_transcript_from_text_and_export(client, tmp_path):
    r = client.post("/transcribe/text", json={"title": "펀딩 회의", "text": MEETING, "auto_todos": True}).json()
    assert r["todos_added"] >= 1 and any("상세페이지" in a["text"] for a in r["actions"])
    assert any(e["start"].endswith("T14:00:00") for e in r["events"])
    t = client.get(f"/transcripts/{r['transcript_id']}").json()
    assert t["summary"] and t["title"] == "펀딩 회의"
    for fmt in ("txt", "md", "srt", "vtt", "docx"):
        out = tmp_path / f"m.{fmt}"
        client.post(f"/transcripts/{t['id']}/export", json={"format": fmt, "out_path": str(out), "masked": True})
        assert out.exists() and out.stat().st_size > 0
    md = (tmp_path / "m.md").read_text(encoding="utf-8")
    assert "## 요약" in md and "## 할 일" in md and "010-****-5678" in md and "010-1234-5678" not in md
    upd = client.put(f"/transcripts/{t['id']}", json={"speakers": {"0": "민수"}, "title": "새 제목"}).json()
    assert upd["title"] == "새 제목"
    assert client.get("/search", params={"q": "슈퍼얼리버드"}).json()["results"][0]["kind"] == "transcript"


def test_srt_timestamps():
    from localbiseo.modules.transcribe.export import to_srt

    srt = to_srt([{"start": 0.5, "end": 3725.25, "text": " 안녕 ", "speaker": 0}], {"0": "민수"})
    assert srt == "1\n00:00:00,500 --> 01:02:05,250\n민수: 안녕\n"


# --- 라이선스 ---
def test_license_trial_and_activation(client, monkeypatch, tmp_path):
    from cryptography.hazmat.primitives import serialization
    from cryptography.hazmat.primitives.asymmetric.ed25519 import Ed25519PrivateKey

    from localbiseo.core import license

    key = Ed25519PrivateKey.generate()
    pub = tmp_path / "pub.pem"
    pub.write_bytes(key.public_key().public_bytes(serialization.Encoding.PEM, serialization.PublicFormat.SubjectPublicKeyInfo))
    monkeypatch.setattr(license, "_pubkey_path", lambda: pub)
    st = client.get("/license").json()
    assert st["status"] == "trial" and st["days_left"] == license.TRIAL_DAYS
    assert client.post("/license", json={"key": "LB1.abc.def"}).status_code == 400
    good = license.sign({"id": "WADIZ-0001", "name": "홍길동", "ed": "pro"}, key)
    r = client.post("/license", json={"key": good[:20] + "\n" + good[20:]})
    assert r.json()["status"] == "licensed" and r.json()["holder"] == "홍길동"
    other = Ed25519PrivateKey.generate()
    forged = license.sign({"id": "X", "name": "해커"}, other)
    assert client.post("/license", json={"key": forged}).status_code == 400


def test_trial_expiry_blocks_core_features(client, app):
    app.state.ctx.db.set_setting("trial_started", (date.today() - timedelta(days=30)).isoformat())
    r = client.post("/mask/files", json={"paths": ["/x.txt"]})
    assert r.status_code == 402
    # 비서 기능은 계속 된다
    assert client.post("/mask/text", json={"text": "010-1234-5678"}).status_code == 200


def test_token_required_when_set(isolated_home):
    from fastapi.testclient import TestClient

    from localbiseo.app import create_app

    with TestClient(create_app(token="secret", data_dir=isolated_home)) as c:
        assert c.get("/health").status_code == 200
        assert c.get("/dashboard").status_code == 401
        assert c.get("/dashboard", headers={"X-Biseo-Token": "secret"}).status_code == 200


def test_dashboard_and_status(client):
    d = client.get("/dashboard").json()
    assert d["todos_open"] == 0 and d["license"]["status"] == "trial"
    s = client.get("/system/status").json()
    assert {c["key"] for c in s["components"]} >= {"stt", "ocr", "rembg", "translate"}


# --- HWP 바이너리 파서 ---
def test_hwp_para_text_controls():
    def wchars(s):
        return s.encode("utf-16-le")

    inline_ctrl = struct.pack("<H", 9) + b"\x00" * 14  # 탭(인라인 컨트롤, 8 WCHAR)
    ext_ctrl = struct.pack("<H", 11) + b"\x00" * 14  # 표/그림(확장 컨트롤, 8 WCHAR)
    data = wchars("이름") + inline_ctrl + wchars("홍길동") + ext_ctrl + wchars("끝") + struct.pack("<H", 13)
    assert _hwp_para_text(data) == "이름\t홍길동끝\n"


def test_hwp_records_parse_sizes():
    payload = "안녕".encode("utf-16-le")
    header = 67 | (0 << 10) | (len(payload) << 20)
    big = b"x" * 5000
    big_header = 50 | (0xFFF << 20)
    stream = struct.pack("<I", header) + payload + struct.pack("<II", big_header, len(big)) + big
    recs = list(_hwp_records(stream))
    assert recs[0] == (67, payload) and recs[1][0] == 50 and len(recs[1][1]) == 5000
