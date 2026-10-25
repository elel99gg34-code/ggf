from __future__ import annotations

import io
import zipfile

import pytest

from localbiseo.modules.masking.detect import detect, mask_text, mask_value, remap_runs
from localbiseo.modules.masking.files import mask_file, mask_hwpx_xml

SAMPLE = (
    "성명: 홍길동, 주민번호 900101-1234567, 연락처 010-1234-5678로 연락 주세요. 사무실 02-345-6789.\n"
    "카드 4111-1111-1111-1111, 국민은행 계좌 123456-01-234567 (예금주 김철수).\n"
    "이메일 gildong.hong@example.com, 차량 12가 3456, 주소 서울특별시 강남구 테헤란로 123, 101동 1001호\n"
    "김민수님께서 방문. 여권 M12345678 면허 11-12-123456-12"
)


def kinds_of(text, **kw):
    return [(f.kind, f.value) for f in detect(text, **kw)]


def test_detects_all_common_kinds():
    found = dict((k, v) for k, v in kinds_of(SAMPLE))
    assert found["rrn"] == "900101-1234567"
    assert found["card"] == "4111-1111-1111-1111"
    assert found["account"] == "123456-01-234567"
    assert found["email"] == "gildong.hong@example.com"
    assert found["plate"] == "12가 3456"
    assert found["passport"] == "M12345678"
    assert found["driver"] == "11-12-123456-12"
    assert found["address"].startswith("서울특별시 강남구 테헤란로 123")
    phones = [v for k, v in kinds_of(SAMPLE) if k == "phone"]
    assert phones == ["010-1234-5678", "02-345-6789"]
    names = [v for k, v in kinds_of(SAMPLE) if k == "name"]
    assert names == ["홍길동", "김철수", "김민수"]


def test_no_false_positives_on_ordinary_text():
    text = "회의는 2026-10-25 오후 3시. 가격은 12,000원이고 수량은 1500개. 선생님께 전달, 고객님 안녕하세요. 마음씨가 곱다."
    assert detect(text) == []


def test_card_requires_luhn():
    assert kinds_of("카드 1234-5678-9012-3456") == []


def test_rrn_without_hyphen_needs_checksum_or_context():
    assert kinds_of("번호 9001011234567") == []  # 체크섬이 맞지 않는 13자리
    assert kinds_of("주민번호 9001011234567")[0][0] == "rrn"


def test_partial_masks_keep_useful_parts():
    masked, _ = mask_text(SAMPLE)
    assert "900101-1******" in masked
    assert "010-****-5678" in masked
    assert "4111-****-****-1111" in masked
    assert "홍*동" in masked
    assert "12가 ****" in masked
    assert "gi**********@example.com" in masked
    assert "서울특별시 강남구 ****" in masked


def test_tag_and_full_modes():
    t = "전화 010-1234-5678"
    assert mask_text(t, mode="tag")[0] == "전화 [전화번호]"
    assert mask_text(t, mode="full")[0] == "전화 ***-****-****"


def test_kinds_filter_does_not_reclassify():
    # 전화번호를 꺼도 계좌로 오인해 가리지 않는다.
    assert detect("010-1234-5678", kinds=["account"]) == []


def test_remap_runs_preserves_run_boundaries():
    runs = ["연락처 010-12", "34-5678 입니다"]
    findings = detect("".join(runs))
    out = remap_runs(runs, findings, "partial")
    assert out == ["연락처 010-**", "**-5678 입니다"]
    out_tag = remap_runs(runs, findings, "tag")
    assert "".join(out_tag) == "연락처 [전화번호] 입니다"


def test_mask_value_name_lengths():
    from localbiseo.modules.masking.detect import Finding

    assert mask_value(Finding("name", 0, 2, "이훈")) == "이*"
    assert mask_value(Finding("name", 0, 4, "남궁민수")) == "남**수"


def test_text_file_cp949(tmp_path):
    src = tmp_path / "고객명단.txt"
    src.write_bytes("이름: 박지성 / 010-9876-5432".encode("cp949"))
    res = mask_file(src, None, "partial")
    out = (tmp_path / "고객명단_마스킹.txt").read_bytes().decode("cp949")
    assert out == "이름: 박*성 / 010-****-5432"
    assert res["by_kind"] == {"name": 1, "phone": 1}


def test_docx_masks_across_runs_and_tables(tmp_path):
    import docx

    doc = docx.Document()
    p = doc.add_paragraph()
    p.add_run("담당자 연락처 010-")
    p.add_run("1234-5678").bold = True
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "주민번호"
    table.cell(0, 1).text = "850505-2345678"
    doc.core_properties.author = "홍길동"
    src = tmp_path / "a.docx"
    doc.save(src)
    res = mask_file(src, None, "partial")
    out = docx.Document(res["dst"])
    runs = out.paragraphs[0].runs
    assert runs[0].text == "담당자 연락처 010-"
    assert runs[1].text == "****-5678" and runs[1].bold
    assert out.tables[0].cell(0, 1).text == "850505-2******"
    assert out.core_properties.author == ""
    assert res["count"] == 2


HWPX_SECTION = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes" ?>'
    '<hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph">'
    '<hp:p id="1"><hp:run charPrIDRef="0"><hp:t>고객 홍길동님 전화 010-</hp:t></hp:run>'
    '<hp:run charPrIDRef="1"><hp:t>1234-5678 &amp; 메일 a@b.co</hp:t></hp:run></hp:p>'
    '<hp:p id="2"><hp:run charPrIDRef="0"><hp:t>010</hp:t></hp:run></hp:p>'
    '<hp:p id="3"><hp:run charPrIDRef="0"><hp:t>-2222-3333<hp:lineBreak/>끝</hp:t></hp:run></hp:p>'
    "</hs:sec>"
)


def test_hwpx_xml_masking_keeps_markup():
    out, findings = mask_hwpx_xml(HWPX_SECTION, None, "partial")
    assert "<hp:t>고객 홍*동님 전화 010-</hp:t>" in out
    assert "<hp:t>****-5678 &amp; 메일 *@b.co</hp:t>" in out
    assert '<hp:run charPrIDRef="1">' in out and "<hp:lineBreak/>" in out
    # 문단을 넘어서는 전화번호(010 / -2222-3333)는 하나로 이어 붙이지 않는다.
    assert {f.kind for f in findings} == {"name", "phone", "email"}
    assert len(findings) == 3


def test_hwpx_file_roundtrip(tmp_path):
    src = tmp_path / "문서.hwpx"
    with zipfile.ZipFile(src, "w") as z:
        z.writestr(zipfile.ZipInfo("mimetype"), "application/hwp+zip")
        z.writestr("Contents/section0.xml", HWPX_SECTION, compress_type=zipfile.ZIP_DEFLATED)
        z.writestr("Preview/PrvText.txt", "고객 홍길동님 전화 010-1234-5678".encode("utf-8"))
        buf = io.BytesIO()
        from PIL import Image

        Image.new("RGB", (40, 60), "red").save(buf, "PNG")
        z.writestr("Preview/PrvImage.png", buf.getvalue())
    res = mask_file(src, None, "partial")
    with zipfile.ZipFile(res["dst"]) as z:
        names = z.namelist()
        assert names[0] == "mimetype"
        assert z.getinfo("mimetype").compress_type == zipfile.ZIP_STORED
        assert "<hp:t>****-5678" in z.read("Contents/section0.xml").decode()
        assert z.read("Preview/PrvText.txt").decode() == "고객 홍*동님 전화 010-****-5678"
        from PIL import Image

        prv = Image.open(io.BytesIO(z.read("Preview/PrvImage.png"))).convert("RGB")
        assert prv.size == (40, 60) and prv.getpixel((5, 5)) == (255, 255, 255)
    assert res["count"] == 3


def _make_pdf(path, font, lines):
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_font("nanum", fname=font)
    pdf.set_font("nanum", size=14)
    for page in lines:
        pdf.add_page()
        for line in page:
            pdf.cell(0, 10, line, new_x="LMARGIN", new_y="NEXT")
    pdf.set_author("작성자 홍길동")
    pdf.output(str(path))


def test_pdf_masks_only_pages_with_pii(tmp_path, korean_font):
    import pypdfium2 as pdfium

    from localbiseo.core.extract import pdf_text

    src = tmp_path / "계약서.pdf"
    _make_pdf(src, korean_font, [["계약자: 홍길동", "연락처 010-1234-5678"], ["두 번째 쪽에는 개인정보가 없어요."]])
    res = mask_file(src, None, "partial")
    assert res["by_kind"] == {"name": 1, "phone": 1}
    text = pdf_text(res["dst"])
    assert "010-1234-5678" not in text and "홍길동" not in text
    assert "두 번째 쪽에는" in text  # 개인정보 없는 쪽은 글자 그대로
    doc = pdfium.PdfDocument(res["dst"])
    assert len(doc) == 2
    assert doc.get_metadata_value("Author") == ""
    # 가린 쪽은 그림이고, 전화번호 위치가 검게 칠해졌는지 확인
    img = doc[0].render(scale=1).to_pil().convert("L")
    assert min(img.getdata()) < 30
    doc.close()


@pytest.mark.parametrize("bad", ["", "abc"])
def test_api_rejects_bad_mode(client, bad):
    r = client.post("/mask/text", json={"text": "x", "mode": bad})
    assert r.status_code == 400


def test_api_mask_text_and_files(client, tmp_path, wait_job):
    r = client.post("/mask/text", json={"text": "전화 010-1234-5678", "mode": "partial"})
    assert r.json()["masked"] == "전화 010-****-5678"
    f = tmp_path / "a.txt"
    f.write_text("메일 test@example.com", encoding="utf-8")
    job = wait_job(client.post("/mask/files", json={"paths": [str(f), str(tmp_path / "없음.txt")]}).json())
    assert job["status"] == "done"
    results = job["result"]["results"]
    assert results[0]["count"] == 1 and "error" in results[1]
