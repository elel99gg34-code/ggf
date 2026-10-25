"""영수증 → 경비 장부(엑셀), 명함 → 연락처(vCard)."""
from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

from fastapi import APIRouter
from pydantic import BaseModel

from ..core import search
from ..core.config import sub_dir
from ..core.context import Ctx
from ..core.db import now_iso
from ..core.errors import AppError
from ..core.korean import looks_like_name
from .screen_ocr import read_image

router = APIRouter(tags=["영수증·명함"])

# --- 영수증 ---
TOTAL_KEYS = ["총결제금액", "결제금액", "받을금액", "받은금액", "합계금액", "총합계", "총액", "합계", "판매총액", "청구금액", "승인금액", "결제 금액", "합 계", "총 액"]
AMOUNT_RE = re.compile(r"(?<![\d,])(\d{1,3}(?:,\d{3})+|\d{3,9})(?:\s*원)?(?![\d,])")
DATE_RES = [
    re.compile(r"(20\d{2})\s*[-./년]\s*(\d{1,2})\s*[-./월]\s*(\d{1,2})"),
    re.compile(r"(?<!\d)(\d{2})[-./](\d{2})[-./](\d{2})(?!\d)"),
]
BIZNO_RE = re.compile(r"(?<!\d)(\d{3})-?(\d{2})-?(\d{5})(?!\d)")
CATEGORY_HINTS = {
    "식비": ["식당", "카페", "커피", "베이커리", "치킨", "피자", "김밥", "국밥", "스타벅스", "투썸", "이디야", "맥도날드", "버거", "분식", "한식", "중식", "일식"],
    "교통": ["주유", "택시", "주차", "하이패스", "고속", "철도", "코레일", "KTX", "SRT", "버스", "오일"],
    "사무용품": ["문구", "오피스", "다이소", "알파", "모닝글로리", "인쇄", "복사"],
    "통신": ["통신", "KT", "SKT", "LG U+", "유플러스"],
    "숙박": ["호텔", "모텔", "펜션", "리조트", "숙박"],
    "마트": ["마트", "이마트", "홈플러스", "롯데마트", "편의점", "GS25", "CU", "세븐일레븐", "이마트24"],
}


def _to_int(s: str) -> int:
    return int(s.replace(",", ""))


def parse_receipt(text: str) -> dict:
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    total = None
    for key in TOTAL_KEYS:
        k = key.replace(" ", "")
        for i, line in enumerate(lines):
            if k in line.replace(" ", ""):
                after = line.replace(" ", "").split(k, 1)[1]
                nums = AMOUNT_RE.findall(after) or (AMOUNT_RE.findall(lines[i + 1]) if i + 1 < len(lines) else [])
                if nums:
                    total = max(_to_int(n) for n in nums)
                    break
        if total:
            break
    if total is None:
        nums = [_to_int(n) for n in AMOUNT_RE.findall(text) if "," in n or n.endswith("00")]
        total = max(nums) if nums else None
    day = None
    for rx in DATE_RES:
        m = rx.search(text)
        if m:
            y, mo, d = m.groups()
            y = int(y) + (2000 if len(y) == 2 else 0)
            try:
                day = date(y, int(mo), int(d)).isoformat()
                break
            except ValueError:
                continue
    biz = BIZNO_RE.search(text)
    merchant = ""
    for line in lines[:5]:
        if re.search(r"[가-힣A-Za-z]{2,}", line) and not re.search(r"영수증|receipt|매출전표|카드|사업자|대표|전화|TEL|주소", line, re.I):
            merchant = re.sub(r"\s{2,}", " ", line)[:40]
            break
    category = ""
    compact = re.sub(r"\s", "", text).lower()
    for cat, words in CATEGORY_HINTS.items():
        if any(w.lower().replace(" ", "") in compact for w in words):
            category = cat
            break
    return {
        "merchant": merchant,
        "date": day,
        "amount": total,
        "biz_no": "-".join(biz.groups()) if biz else "",
        "category": category,
    }


# --- 명함 ---
JOB_TITLES = ["대표이사", "대표", "부사장", "전무", "상무", "이사", "본부장", "실장", "부장", "차장", "과장", "대리", "주임", "사원", "팀장", "센터장", "원장", "교수", "연구원", "책임", "선임", "매니저", "디자이너", "개발자", "변호사", "세무사", "노무사", "회계사", "CEO", "CTO", "CFO", "Manager", "Director"]
COMPANY_RE = re.compile(r"(?:주식회사|\(주\)|㈜|유한회사|\(유\)|재단법인|사단법인|법무법인|세무법인|회계법인|Co\.,?\s?Ltd|Inc\.?|Corp\.?)", re.I)
EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
WEB_RE = re.compile(r"(?:https?://)?(?:www\.)[A-Za-z0-9.-]+\.[A-Za-z]{2,}(?:/\S*)?|https?://[A-Za-z0-9.-]+\.[A-Za-z]{2,}(?:/\S*)?")
PHONE_RE = re.compile(r"(?:\+82[-.\s]?)?0?\d{1,2}[-.\s)]?\d{3,4}[-.\s]?\d{4}")
ADDR_RE = re.compile(r"(서울|부산|대구|인천|광주|대전|울산|세종|경기|강원|충북|충남|전북|전남|경북|경남|제주|충청|전라|경상)\S*\s.+")


def _norm_phone(s: str) -> str:
    d = re.sub(r"\D", "", s)
    if d.startswith("82"):
        d = "0" + d[2:].lstrip("0")
    if d.startswith("02"):
        return f"02-{d[2:-4]}-{d[-4:]}"
    return f"{d[:3]}-{d[3:-4]}-{d[-4:]}" if len(d) >= 9 else s


def parse_card(text: str) -> dict:
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    out = {k: "" for k in ("name", "company", "title", "mobile", "phone", "fax", "email", "address", "website")}
    for line in lines:
        low = line.lower()
        if not out["email"] and (m := EMAIL_RE.search(line)):
            out["email"] = m.group(0)
            continue
        if not out["website"] and (m := WEB_RE.search(line)) and "@" not in line:
            out["website"] = m.group(0)
            continue
        phones = PHONE_RE.findall(line)
        if phones:
            for p in phones:
                np_ = _norm_phone(p)
                before = line[: line.find(p)].lower()
                if re.search(r"fax|팩스|f\.", before) and not out["fax"]:
                    out["fax"] = np_
                elif np_.startswith("01") and not out["mobile"]:
                    out["mobile"] = np_
                elif not out["phone"]:
                    out["phone"] = np_
            continue
        if not out["address"] and ADDR_RE.match(line):
            out["address"] = line
            continue
        if not out["company"] and COMPANY_RE.search(line):
            out["company"] = line
            continue
        if "fax" in low or "tel" in low:
            continue
        if not out["title"]:
            for t in JOB_TITLES:
                if t in line:
                    out["title"] = t
                    rest = line.replace(t, "").strip(" /|·")
                    if not out["name"] and looks_like_name(rest.replace(" ", "")):
                        out["name"] = rest.replace(" ", "")
                    break
            if out["title"]:
                continue
        compact = line.replace(" ", "")
        if not out["name"] and looks_like_name(compact):
            out["name"] = compact
    return out


def to_vcard(c: dict) -> str:
    lines = ["BEGIN:VCARD", "VERSION:3.0", f"FN:{c.get('name', '')}"]
    name = c.get("name", "")
    if name:
        lines.append(f"N:{name[0]};{name[1:]};;;")
    if c.get("company"):
        lines.append(f"ORG:{c['company']}")
    if c.get("title"):
        lines.append(f"TITLE:{c['title']}")
    for key, typ in (("mobile", "CELL"), ("phone", "WORK"), ("fax", "FAX")):
        if c.get(key):
            lines.append(f"TEL;TYPE={typ}:{c[key]}")
    if c.get("email"):
        lines.append(f"EMAIL;TYPE=WORK:{c['email']}")
    if c.get("address"):
        lines.append(f"ADR;TYPE=WORK:;;{c['address']};;;;")
    if c.get("website"):
        lines.append(f"URL:{c['website']}")
    lines.append("END:VCARD")
    return "\r\n".join(lines) + "\r\n"


class ParseIn(BaseModel):
    text: str | None = None
    path: str | None = None


class ExpenseIn(BaseModel):
    date: str | None = None
    merchant: str = ""
    amount: int | None = None
    category: str = ""
    memo: str = ""
    biz_no: str = ""


class ContactIn(BaseModel):
    name: str = ""
    company: str = ""
    title: str = ""
    mobile: str = ""
    phone: str = ""
    fax: str = ""
    email: str = ""
    address: str = ""
    website: str = ""


class ExportIn(BaseModel):
    month: str | None = None  # YYYY-MM
    out_path: str | None = None


def _source_text(body: ParseIn) -> str:
    if body.text:
        return body.text
    if body.path:
        return read_image(body.path)["text"]
    raise AppError("글자나 이미지 경로를 넣어 주세요.")


@router.post("/receipts/parse")
def receipt_parse(body: ParseIn):
    text = _source_text(body)
    return {"text": text, "fields": parse_receipt(text)}


@router.get("/expenses")
def list_expenses(ctx: Ctx, month: str | None = None):
    if month:
        return ctx.db.query("SELECT * FROM expenses WHERE date LIKE ? ORDER BY date DESC, id DESC", (f"{month}%",))
    return ctx.db.query("SELECT * FROM expenses ORDER BY date DESC, id DESC")


@router.post("/expenses")
def add_expense(body: ExpenseIn, ctx: Ctx):
    eid = ctx.db.insert(
        "INSERT INTO expenses(date, merchant, amount, category, memo, biz_no, created_at) VALUES(?,?,?,?,?,?,?)",
        (body.date, body.merchant, body.amount, body.category, body.memo, body.biz_no, now_iso()),
    )
    return ctx.db.one("SELECT * FROM expenses WHERE id=?", (eid,))


@router.delete("/expenses/{eid}")
def delete_expense(eid: int, ctx: Ctx):
    ctx.db.execute("DELETE FROM expenses WHERE id=?", (eid,))
    return {"ok": True}


@router.post("/expenses/export")
def export_expenses(body: ExportIn, ctx: Ctx):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    rows = list_expenses(ctx, body.month)
    wb = Workbook()
    ws = wb.active
    ws.title = "경비"
    head = ["날짜", "가맹점", "금액", "분류", "메모", "사업자번호"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8EEF9")
    for r in rows:
        ws.append([r["date"], r["merchant"], r["amount"], r["category"], r["memo"], r["biz_no"]])
    total_row = len(rows) + 2
    ws.cell(row=total_row, column=2, value="합계").font = Font(bold=True)
    ws.cell(row=total_row, column=3, value=f"=SUM(C2:C{total_row - 1})").font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        for c in row:
            c.number_format = "#,##0"
    for col, width in zip("ABCDEF", (12, 28, 12, 10, 30, 14)):
        ws.column_dimensions[col].width = width
    path = Path(body.out_path) if body.out_path else sub_dir("exports") / f"경비_{body.month or datetime.now().strftime('%Y%m%d')}.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return {"path": str(path), "count": len(rows), "total": sum(r["amount"] or 0 for r in rows)}


@router.post("/cards/parse")
def card_parse(body: ParseIn):
    text = _source_text(body)
    return {"text": text, "fields": parse_card(text)}


@router.get("/contacts")
def list_contacts(ctx: Ctx, q: str = ""):
    if q:
        like = f"%{q}%"
        return ctx.db.query("SELECT * FROM contacts WHERE name LIKE ? OR company LIKE ? OR email LIKE ? ORDER BY name", (like, like, like))
    return ctx.db.query("SELECT * FROM contacts ORDER BY id DESC")


@router.post("/contacts")
def add_contact(body: ContactIn, ctx: Ctx):
    data = body.model_dump()
    cid = ctx.db.insert(
        f"INSERT INTO contacts({', '.join(data)}, created_at) VALUES({', '.join('?' * len(data))}, ?)",
        (*data.values(), now_iso()),
    )
    search.index(ctx.db, "contact", cid, body.name or body.company, " ".join(v for v in data.values() if v))
    return ctx.db.one("SELECT * FROM contacts WHERE id=?", (cid,))


@router.delete("/contacts/{cid}")
def delete_contact(cid: int, ctx: Ctx):
    ctx.db.execute("DELETE FROM contacts WHERE id=?", (cid,))
    search.remove(ctx.db, "contact", cid)
    return {"ok": True}


@router.post("/contacts/export")
def export_contacts(body: ExportIn, ctx: Ctx):
    rows = ctx.db.query("SELECT * FROM contacts ORDER BY id")
    path = Path(body.out_path) if body.out_path else sub_dir("exports") / f"연락처_{datetime.now():%Y%m%d}.vcf"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("".join(to_vcard(r) for r in rows), encoding="utf-8", newline="")
    return {"path": str(path), "count": len(rows)}
