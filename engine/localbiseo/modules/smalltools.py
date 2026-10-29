"""작은 도구 10가지: 글자 수, 금액 한글·부가세, 날짜 계산, QR, 비밀번호, 한영 타자 변환, 글 정리, 단위 변환, 번호 확인, 표 합치기."""
from __future__ import annotations

import base64
import csv
import io
import math
import re
import secrets
import string
from datetime import date, datetime, timedelta
from pathlib import Path

from fastapi import APIRouter
from pydantic import BaseModel, Field

from ..core.config import output_path, sub_dir
from ..core.errors import AppError
from .masking.detect import bizno_checksum_ok, luhn_ok, rrn_checksum_ok

router = APIRouter(prefix="/tools", tags=["작은 도구"])


# --- 1. 글자 수 ---
class TextIn(BaseModel):
    text: str


@router.post("/count")
def count(body: TextIn):
    t = body.text
    no_space = re.sub(r"\s", "", t)
    return {
        "chars": len(t),
        "chars_no_space": len(no_space),
        "words": len(t.split()),
        "lines": t.count("\n") + 1 if t else 0,
        "paragraphs": len([p for p in re.split(r"\n\s*\n", t) if p.strip()]),
        "bytes_utf8": len(t.encode("utf-8")),
        # 채용 사이트·공공기관 입력 칸이 흔히 쓰는 기준: 한글 2바이트, 영문·숫자 1바이트
        "bytes_ko": sum(1 if ord(c) < 128 else 2 for c in t),
        "manuscript_pages": math.ceil(len(t) / 200) if t else 0,  # 200자 원고지 (대략)
        "reading_minutes": round(len(no_space) / 500, 1),  # 1분에 500자 정도 읽는다고 보고
    }


# --- 2. 금액 한글·부가세 ---
DIGITS = "영일이삼사오육칠팔구"
SMALL_UNITS = ["", "십", "백", "천"]
BIG_UNITS = ["", "만", "억", "조", "경"]


def _groups(n: int) -> list[int]:
    out = []
    while n:
        out.append(n % 10000)
        n //= 10000
    return out  # 낮은 자리부터


def _group_hangul(g: int, formal: bool) -> str:
    s = ""
    for p in range(3, -1, -1):
        d = g // 10**p % 10
        if d:
            s += ("" if d == 1 and p > 0 and not formal else DIGITS[d]) + SMALL_UNITS[p]
    return s


def hangul_amount(n: int, formal: bool = False) -> str:
    """formal: 일백이십삼만사천오백육십칠 (수표·계약서) / 아니면: 백이십삼만 사천오백육십칠 (읽는 말)"""
    if n == 0:
        return "영"
    parts = []
    for i, g in reversed(list(enumerate(_groups(n)))):
        if not g:
            continue
        word = _group_hangul(g, formal)
        if not formal and g == 1 and i == 1:
            word = ""  # 말할 때는 '일만'이 아니라 '만' (억·조는 '일억', '일조')
        parts.append(word + BIG_UNITS[i])
    return ("" if formal else " ").join(parts)


def short_amount(n: int) -> str:
    """1234567 → 123만 4,567원"""
    if n == 0:
        return "0원"
    parts = []
    for i, g in reversed(list(enumerate(_groups(n)))):
        if g:
            parts.append(f"{g:,}{BIG_UNITS[i]}")
    return " ".join(parts) + "원"


class MoneyIn(BaseModel):
    amount: str
    vat: str = "none"  # none | add(공급가에 부가세 더하기) | split(합계에서 나누기)


@router.post("/money")
def money(body: MoneyIn):
    raw = re.sub(r"[,\s원₩]", "", body.amount)
    if not raw.isdigit():
        raise AppError("금액은 숫자로 넣어 주세요. (예: 1,234,000)")
    n = int(raw)
    if n >= 10**16:
        raise AppError("1경 원보다 작은 금액만 계산할 수 있어요.")
    out = {"amount": n, "formatted": f"{n:,}원", "short": short_amount(n), "reading": hangul_amount(n) + " 원",
           "formal": f"일금 {hangul_amount(n, formal=True)}원정"}
    if body.vat == "add":
        vat = n // 10  # 원 단위 미만은 버린다
        out["vat"] = {"supply": n, "tax": vat, "total": n + vat}
    elif body.vat == "split":
        supply = (n * 10 + 5) // 11  # 합계 ÷ 1.1, 반올림
        out["vat"] = {"supply": supply, "tax": n - supply, "total": n}
    return out


# --- 3. 날짜 계산 ---
# 관공서 공휴일 (대체공휴일 포함). 임시공휴일은 정해지면 새 버전에 넣는다.
HOLIDAYS: dict[str, str] = {
    # 2025
    "2025-01-01": "신정", "2025-01-27": "임시공휴일", "2025-01-28": "설날 연휴", "2025-01-29": "설날", "2025-01-30": "설날 연휴",
    "2025-03-01": "삼일절", "2025-03-03": "대체공휴일(삼일절)", "2025-05-05": "어린이날·부처님오신날", "2025-05-06": "대체공휴일",
    "2025-06-03": "대통령 선거일", "2025-06-06": "현충일", "2025-08-15": "광복절", "2025-10-03": "개천절", "2025-10-05": "추석 연휴",
    "2025-10-06": "추석", "2025-10-07": "추석 연휴", "2025-10-08": "대체공휴일(추석)", "2025-10-09": "한글날", "2025-12-25": "성탄절",
    # 2026 (노동절·제헌절이 공휴일이 됐다)
    "2026-01-01": "신정", "2026-02-16": "설날 연휴", "2026-02-17": "설날", "2026-02-18": "설날 연휴", "2026-03-01": "삼일절",
    "2026-03-02": "대체공휴일(삼일절)", "2026-05-01": "노동절", "2026-05-05": "어린이날", "2026-05-24": "부처님오신날",
    "2026-05-25": "대체공휴일(부처님오신날)", "2026-06-03": "전국동시지방선거일", "2026-06-06": "현충일", "2026-07-17": "제헌절",
    "2026-08-15": "광복절", "2026-08-17": "대체공휴일(광복절)", "2026-09-24": "추석 연휴", "2026-09-25": "추석", "2026-09-26": "추석 연휴",
    "2026-10-03": "개천절", "2026-10-05": "대체공휴일(개천절)", "2026-10-09": "한글날", "2026-12-25": "성탄절",
    # 2027
    "2027-01-01": "신정", "2027-02-06": "설날 연휴", "2027-02-07": "설날", "2027-02-08": "설날 연휴", "2027-02-09": "대체공휴일(설날)",
    "2027-03-01": "삼일절", "2027-05-01": "노동절", "2027-05-03": "대체공휴일(노동절)", "2027-05-05": "어린이날",
    "2027-05-13": "부처님오신날", "2027-06-06": "현충일", "2027-07-17": "제헌절", "2027-07-19": "대체공휴일(제헌절)",
    "2027-08-15": "광복절", "2027-08-16": "대체공휴일(광복절)", "2027-09-14": "추석 연휴", "2027-09-15": "추석", "2027-09-16": "추석 연휴",
    "2027-10-03": "개천절", "2027-10-04": "대체공휴일(개천절)", "2027-10-09": "한글날", "2027-10-11": "대체공휴일(한글날)",
    "2027-12-25": "성탄절", "2027-12-27": "대체공휴일(성탄절)",
}
HOLIDAY_YEARS = (2025, 2027)
WEEKDAYS = "월화수목금토일"


def is_workday(d: date) -> bool:
    return d.weekday() < 5 and d.isoformat() not in HOLIDAYS


def _day(d: date) -> dict:
    return {"date": d.isoformat(), "weekday": WEEKDAYS[d.weekday()], "holiday": HOLIDAYS.get(d.isoformat())}


def _parse_date(s: str | None, name: str) -> date:
    if not s:
        raise AppError(f"{name}을 넣어 주세요.")
    try:
        return date.fromisoformat(s.strip().replace(".", "-").replace("/", "-"))
    except ValueError as e:
        raise AppError(f"{name}은 2026-10-29처럼 넣어 주세요.") from e


def add_workdays(start: date, n: int) -> date:
    step = 1 if n >= 0 else -1
    d, left = start, abs(n)
    while left:
        d += timedelta(days=step)
        if is_workday(d):
            left -= 1
    return d


class DateIn(BaseModel):
    mode: str = "between"  # between | add | dday
    start: str | None = None
    end: str | None = None
    days: int = 0
    workdays: bool = False  # add에서 주말·공휴일 빼고 세기
    today: str | None = None  # 테스트용


@router.post("/date")
def date_calc(body: DateIn):
    today = date.fromisoformat(body.today) if body.today else date.today()
    notes = []
    if body.mode == "between":
        a, b = _parse_date(body.start, "시작일"), _parse_date(body.end, "끝날")
        lo, hi = min(a, b), max(a, b)
        days = [lo + timedelta(days=i) for i in range((hi - lo).days + 1)]
        out = {
            "start": _day(a), "end": _day(b), "days": (b - a).days, "days_inclusive": len(days),
            "weeks": f"{(hi - lo).days // 7}주 {(hi - lo).days % 7}일",
            "workdays": sum(is_workday(d) for d in days),
            "holidays": [_day(d) for d in days if d.isoformat() in HOLIDAYS],
        }
        years = {lo.year, hi.year}
    elif body.mode == "add":
        a = _parse_date(body.start, "시작일") if body.start else today
        res = add_workdays(a, body.days) if body.workdays else a + timedelta(days=body.days)
        out = {"start": _day(a), "result": _day(res), "days": body.days, "workdays": body.workdays}
        years = {a.year, res.year}
    elif body.mode == "dday":
        t = _parse_date(body.end or body.start, "목표일")
        diff = (t - today).days
        label = "D-day" if diff == 0 else (f"D-{diff}" if diff > 0 else f"D+{-diff}")
        lo, hi = min(today, t), max(today, t)
        between = [lo + timedelta(days=i) for i in range(1, (hi - lo).days + 1)]
        out = {"target": _day(t), "today": _day(today), "days": diff, "label": label, "workdays": sum(is_workday(d) for d in between)}
        years = {lo.year, hi.year}
    else:
        raise AppError("계산 방식은 between, add, dday 중 하나예요.")
    if any(y < HOLIDAY_YEARS[0] or y > HOLIDAY_YEARS[1] for y in years):
        notes.append(f"공휴일 표는 {HOLIDAY_YEARS[0]}~{HOLIDAY_YEARS[1]}년만 들어 있어요. 그 밖의 해는 주말만 빼고 셌어요.")
    out["notes"] = notes
    return out


@router.get("/holidays")
def holidays(year: int):
    return [{**_day(date.fromisoformat(d)), "name": n} for d, n in HOLIDAYS.items() if d.startswith(f"{year}-")]


# --- 4. QR 코드 ---
def _wifi_escape(s: str) -> str:
    return re.sub(r'([\\;,:"])', r"\\\1", s)


class QrIn(BaseModel):
    kind: str = "text"  # text | wifi
    text: str = ""
    ssid: str = ""
    password: str = ""
    security: str = "WPA"  # WPA | WEP | nopass
    scale: int = Field(10, ge=2, le=40)
    out_path: str | None = None


@router.post("/qr")
def qr(body: QrIn):
    import segno

    if body.kind == "wifi":
        if not body.ssid.strip():
            raise AppError("와이파이 이름(SSID)을 넣어 주세요.")
        sec = body.security if body.security in ("WPA", "WEP", "nopass") else "WPA"
        content = f"WIFI:T:{sec};S:{_wifi_escape(body.ssid)};" + (f"P:{_wifi_escape(body.password)};" if sec != "nopass" else "") + ";"
        stem = f"QR_와이파이_{body.ssid}"
    else:
        content = body.text.strip()
        if not content:
            raise AppError("QR에 담을 글이나 주소를 넣어 주세요.")
        stem = "QR_" + re.sub(r"^https?://(www\.)?", "", content)[:24]
    try:
        code = segno.make(content, error="m", micro=False)
    except Exception as e:  # noqa: BLE001 - 너무 긴 글
        raise AppError("글이 너무 길어서 QR 코드에 담을 수 없어요.") from e
    buf = io.BytesIO()
    code.save(buf, kind="png", scale=body.scale, border=2)
    if body.out_path:
        dst = Path(body.out_path)
        dst.parent.mkdir(parents=True, exist_ok=True)
    else:
        safe = re.sub(r'[\\/:*?"<>|\s]+', "_", stem).strip("_") or "QR"
        dst = output_path(Path(f"{safe}.png"), "", out_dir=sub_dir("exports"))
    dst.write_bytes(buf.getvalue())
    return {"dst": str(dst), "content": content, "version": code.version,
            "data_url": "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()}


# --- 5. 비밀번호 ---
AMBIGUOUS = set("Il1O0o|`'\"")
SYMBOLS = "!@#$%^&*-_=+?"


class PasswordIn(BaseModel):
    length: int = Field(16, ge=6, le=64)
    count: int = Field(5, ge=1, le=20)
    upper: bool = True
    digits: bool = True
    symbols: bool = True
    avoid_ambiguous: bool = True


@router.post("/password")
def password(body: PasswordIn):
    pools = [string.ascii_lowercase]
    if body.upper:
        pools.append(string.ascii_uppercase)
    if body.digits:
        pools.append(string.digits)
    if body.symbols:
        pools.append(SYMBOLS)
    if body.avoid_ambiguous:
        pools = ["".join(c for c in p if c not in AMBIGUOUS) for p in pools]
    if body.length < len(pools):
        raise AppError("길이가 너무 짧아요.")
    alphabet = "".join(pools)
    out = []
    for _ in range(body.count):
        # 고른 종류가 한 글자씩은 꼭 들어가게 한다.
        chars = [secrets.choice(p) for p in pools] + [secrets.choice(alphabet) for _ in range(body.length - len(pools))]
        for i in range(len(chars) - 1, 0, -1):
            j = secrets.randbelow(i + 1)
            chars[i], chars[j] = chars[j], chars[i]
        out.append("".join(chars))
    bits = round(body.length * math.log2(len(alphabet)))
    strength = "약함" if bits < 50 else ("보통" if bits < 80 else "강함")
    return {"passwords": out, "bits": bits, "strength": strength}


# --- 6. 한영 타자 변환 (두벌식) ---
CHO = "ㄱㄲㄴㄷㄸㄹㅁㅂㅃㅅㅆㅇㅈㅉㅊㅋㅌㅍㅎ"
JUNG = "ㅏㅐㅑㅒㅓㅔㅕㅖㅗㅘㅙㅚㅛㅜㅝㅞㅟㅠㅡㅢㅣ"
JONG = ["", "ㄱ", "ㄲ", "ㄳ", "ㄴ", "ㄵ", "ㄶ", "ㄷ", "ㄹ", "ㄺ", "ㄻ", "ㄼ", "ㄽ", "ㄾ", "ㄿ", "ㅀ", "ㅁ", "ㅂ", "ㅄ", "ㅅ", "ㅆ",
        "ㅇ", "ㅈ", "ㅊ", "ㅋ", "ㅌ", "ㅍ", "ㅎ"]
KEY_TO_JAMO = dict(zip("qwertyuiopasdfghjklzxcvbnm", "ㅂㅈㄷㄱㅅㅛㅕㅑㅐㅔㅁㄴㅇㄹㅎㅗㅓㅏㅣㅋㅌㅊㅍㅠㅜㅡ"))
KEY_TO_JAMO.update({"Q": "ㅃ", "W": "ㅉ", "E": "ㄸ", "R": "ㄲ", "T": "ㅆ", "O": "ㅒ", "P": "ㅖ"})
VOWEL_PAIRS = {("ㅗ", "ㅏ"): "ㅘ", ("ㅗ", "ㅐ"): "ㅙ", ("ㅗ", "ㅣ"): "ㅚ", ("ㅜ", "ㅓ"): "ㅝ", ("ㅜ", "ㅔ"): "ㅞ", ("ㅜ", "ㅣ"): "ㅟ", ("ㅡ", "ㅣ"): "ㅢ"}
FINAL_PAIRS = {("ㄱ", "ㅅ"): "ㄳ", ("ㄴ", "ㅈ"): "ㄵ", ("ㄴ", "ㅎ"): "ㄶ", ("ㄹ", "ㄱ"): "ㄺ", ("ㄹ", "ㅁ"): "ㄻ", ("ㄹ", "ㅂ"): "ㄼ",
               ("ㄹ", "ㅅ"): "ㄽ", ("ㄹ", "ㅌ"): "ㄾ", ("ㄹ", "ㅍ"): "ㄿ", ("ㄹ", "ㅎ"): "ㅀ", ("ㅂ", "ㅅ"): "ㅄ"}
SPLIT_FINAL = {v: k for k, v in FINAL_PAIRS.items()}
SPLIT_VOWEL = {v: k for k, v in VOWEL_PAIRS.items()}
JAMO_TO_KEY = {v: k for k, v in KEY_TO_JAMO.items()}


def _syllable(cho: str, jung: str, jong: str = "") -> str:
    return chr(0xAC00 + (CHO.index(cho) * 21 + JUNG.index(jung)) * 28 + JONG.index(jong))


def en_to_ko(text: str) -> str:
    out: list[str] = []
    cho = jung = jong = ""

    def flush():
        nonlocal cho, jung, jong
        if cho and jung:
            out.append(_syllable(cho, jung, jong))
        else:
            out.append(cho + jung)
        cho = jung = jong = ""

    for ch in text:
        j = KEY_TO_JAMO.get(ch) or KEY_TO_JAMO.get(ch.lower())
        if not j:
            flush()
            out.append(ch)
            continue
        if j in CHO:  # 자음
            if not jung:
                flush()
                cho = j
            elif not cho:
                flush()
                cho = j
            elif not jong:
                if j in JONG:
                    jong = j
                else:
                    flush()
                    cho = j
            elif (jong, j) in FINAL_PAIRS:
                jong = FINAL_PAIRS[(jong, j)]
            else:
                flush()
                cho = j
        else:  # 모음
            if jong:
                if jong in SPLIT_FINAL:
                    first, second = SPLIT_FINAL[jong]
                    jong = first
                    flush()
                    cho = second
                else:
                    moved = jong
                    jong = ""
                    flush()
                    cho = moved
                jung = j
            elif jung:
                if (jung, j) in VOWEL_PAIRS:
                    jung = VOWEL_PAIRS[(jung, j)]
                else:
                    flush()
                    jung = j
            else:
                jung = j
    flush()
    return "".join(out)


def ko_to_en(text: str) -> str:
    out = []
    for ch in text:
        code = ord(ch) - 0xAC00
        if 0 <= code < 11172:
            parts = [CHO[code // 588], JUNG[code % 588 // 28], JONG[code % 28]]
        elif ch in JAMO_TO_KEY or ch in SPLIT_FINAL or ch in SPLIT_VOWEL:
            parts = [ch]
        else:
            out.append(ch)
            continue
        for p in parts:
            for piece in SPLIT_FINAL.get(p) or SPLIT_VOWEL.get(p) or (p,):
                out.append(JAMO_TO_KEY.get(piece, ""))
    return "".join(out)


class KeyboardIn(BaseModel):
    text: str
    direction: str = "auto"  # auto | en2ko | ko2en


@router.post("/keyboard")
def keyboard(body: KeyboardIn):
    direction = body.direction
    if direction == "auto":
        direction = "ko2en" if re.search(r"[가-힣ㄱ-ㅣ]", body.text) else "en2ko"
    if direction not in ("en2ko", "ko2en"):
        raise AppError("방향은 auto, en2ko, ko2en 중 하나예요.")
    return {"direction": direction, "text": en_to_ko(body.text) if direction == "en2ko" else ko_to_en(body.text)}


# --- 7. 글 정리 ---
CLEAN_OPS = {
    "join_lines": "줄바꿈 합치기 (PDF에서 복사한 글)",
    "bullets": "앞머리 기호·번호 지우기",
    "halfwidth": "전각 글자를 반각으로",
    "emoji": "이모지 지우기",
    "trim": "줄 앞뒤 공백 지우기",
    "spaces": "겹친 띄어쓰기 하나로",
    "blank": "빈 줄 지우기",
    "dedupe": "같은 줄 지우기",
    "sort": "가나다순 정렬",
    "number": "줄 번호 붙이기",
}
EMOJI_RE = re.compile("[\U0001F000-\U0001FAFF☀-➿️‍]")
BULLET_RE = re.compile(r"^\s*(?:[-*•·▪◦●○■□▶►✓✔]|\d{1,3}[.)]|[①-⑳]|[가-하][.)])\s*")


class CleanIn(BaseModel):
    text: str
    ops: list[str] = Field(default_factory=lambda: ["trim", "spaces"])


@router.post("/clean")
def clean(body: CleanIn):
    unknown = set(body.ops) - set(CLEAN_OPS)
    if unknown:
        raise AppError(f"알 수 없는 정리 방법: {', '.join(sorted(unknown))}")
    t = body.text.replace("\r\n", "\n")
    ops = set(body.ops)
    if "join_lines" in ops:
        paras = re.split(r"\n\s*\n", t)
        t = "\n\n".join(re.sub(r"\s*\n\s*", " ", p.strip()) for p in paras)
    lines = t.split("\n")
    if "bullets" in ops:
        lines = [BULLET_RE.sub("", l) for l in lines]
    if "halfwidth" in ops:
        lines = [l.translate({**{c: c - 0xFEE0 for c in range(0xFF01, 0xFF5F)}, 0x3000: 0x20}) for l in lines]
    if "emoji" in ops:
        lines = [EMOJI_RE.sub("", l) for l in lines]
    if "trim" in ops:
        lines = [l.strip() for l in lines]
    if "spaces" in ops:
        lines = [re.sub(r"[ \t ]{2,}", " ", l) for l in lines]
    if "blank" in ops:
        lines = [l for l in lines if l.strip()]
    if "dedupe" in ops:
        seen: set[str] = set()
        lines = [l for l in lines if not (l in seen or seen.add(l)) or not l.strip()]
    if "sort" in ops:
        lines = sorted(lines, key=lambda l: (not l.strip(), l))
    if "number" in ops:
        n = 0
        numbered = []
        for l in lines:
            if l.strip():
                n += 1
                numbered.append(f"{n}. {l}")
            else:
                numbered.append(l)
        lines = numbered
    out = "\n".join(lines)
    return {"text": out, "before": len(body.text), "after": len(out)}


@router.get("/clean/ops")
def clean_ops():
    return [{"key": k, "label": v} for k, v in CLEAN_OPS.items()]


# --- 8. 단위 변환 ---
UNITS: dict[str, dict[str, float]] = {
    "넓이": {"㎡": 1, "평": 400 / 121, "ft²": 0.09290304, "a": 100, "ha": 10000},
    "길이": {"mm": 0.001, "cm": 0.01, "m": 1, "km": 1000, "in": 0.0254, "ft": 0.3048, "yd": 0.9144, "mi": 1609.344, "자": 10 / 33},
    "무게": {"g": 1, "kg": 1000, "톤": 1_000_000, "근": 600, "돈": 3.75, "냥": 37.5, "lb": 453.59237, "oz": 28.349523125},
    "부피": {"mL": 1, "L": 1000, "되": 1803.9, "말": 18039, "gal": 3785.411784, "컵": 200},
    "온도": {"°C": 1, "°F": 1, "K": 1},
}
UNIT_NOTES = {"근": "고기 기준 600g (채소는 375g)", "평": "1평 = 400/121㎡ ≈ 3.3058㎡", "컵": "종이컵 200mL"}


def _to_celsius(v: float, u: str) -> float:
    return {"°C": v, "°F": (v - 32) * 5 / 9, "K": v - 273.15}[u]


def _from_celsius(c: float, u: str) -> float:
    return {"°C": c, "°F": c * 9 / 5 + 32, "K": c + 273.15}[u]


def _nice(x: float) -> float:
    if x == 0:
        return 0
    digits = max(0, 6 - int(math.floor(math.log10(abs(x)))) - 1)
    return round(x, min(digits, 10))


class UnitIn(BaseModel):
    value: float
    unit: str


@router.get("/units")
def unit_list():
    return {cat: list(units) for cat, units in UNITS.items()}


@router.post("/units")
def convert(body: UnitIn):
    cat = next((c for c, units in UNITS.items() if body.unit in units), None)
    if not cat:
        raise AppError("알 수 없는 단위예요.")
    if cat == "온도":
        c = _to_celsius(body.value, body.unit)
        results = [{"unit": u, "value": _nice(_from_celsius(c, u))} for u in UNITS[cat]]
    else:
        base = body.value * UNITS[cat][body.unit]
        results = [{"unit": u, "value": _nice(base / f), "note": UNIT_NOTES.get(u)} for u, f in UNITS[cat].items()]
    return {"category": cat, "results": results}


# --- 9. 번호 확인 ---
BIZ_TYPES = [
    (range(1, 80), "개인 과세사업자"), (range(80, 81), "아파트 관리사무소 등"), (range(81, 82), "영리법인 본점"),
    (range(82, 83), "비영리법인"), (range(83, 84), "국가·지방자치단체"), (range(84, 85), "외국법인"), (range(85, 86), "영리법인 지점"),
    (range(86, 89), "영리법인 본점"), (range(89, 90), "법인 아닌 단체"), (range(90, 100), "개인 면세사업자"),
]
CARD_BRANDS = [(r"^4", "비자"), (r"^(5[1-5]|2[2-7])", "마스터카드"), (r"^3[47]", "아메리칸 익스프레스"), (r"^35", "JCB"),
               (r"^62", "유니온페이"), (r"^9", "국내 전용")]


def corp_no_ok(d: str) -> bool:
    total = sum(int(c) * (1 if i % 2 == 0 else 2) for i, c in enumerate(d[:12]))
    return (10 - total % 10) % 10 == int(d[12])


def _rrn_date_ok(d: str) -> bool:
    century = {"1": 1900, "2": 1900, "3": 2000, "4": 2000, "5": 1900, "6": 1900, "7": 2000, "8": 2000}.get(d[6])
    if not century:
        return False
    try:
        date(century + int(d[:2]), int(d[2:4]), int(d[4:6]))
        return True
    except ValueError:
        return False


class CheckIn(BaseModel):
    value: str


@router.post("/check")
def check(body: CheckIn):
    d = re.sub(r"\D", "", body.value)
    if len(d) == 10 and not d.startswith("0"):
        mid = int(d[3:5])
        kind = next((label for r, label in BIZ_TYPES if mid in r), "")
        ok = bizno_checksum_ok(d)
        return {"kind": "사업자등록번호", "valid": ok, "formatted": f"{d[:3]}-{d[3:5]}-{d[5:]}",
                "detail": kind, "note": "" if ok else "검증 숫자가 맞지 않아요. 잘못 적었을 수 있어요."}
    if len(d) == 13:
        if _rrn_date_ok(d):
            ok = rrn_checksum_ok(d)
            note = "형식과 검증 숫자가 맞아요." if ok else "2020년 10월 이후 받은 번호는 검증 숫자가 없어서 생년월일·성별 자리만 확인했어요."
            return {"kind": "주민등록번호", "valid": True, "checksum": ok, "formatted": f"{d[:6]}-{d[6:]}", "detail": "", "note": note}
        ok = corp_no_ok(d)
        return {"kind": "법인등록번호", "valid": ok, "formatted": f"{d[:6]}-{d[6:]}", "detail": "",
                "note": "" if ok else "검증 숫자가 맞지 않아요."}
    if 14 <= len(d) <= 16 and not d.startswith("0"):
        brand = next((label for rx, label in CARD_BRANDS if re.match(rx, d)), "")
        ok = luhn_ok(d)
        groups = " ".join(d[i : i + 4] for i in range(0, len(d), 4))
        return {"kind": "카드번호", "valid": ok, "formatted": groups, "detail": brand, "note": "" if ok else "검증 숫자가 맞지 않아요."}
    if d.startswith("0") and 9 <= len(d) <= 11:
        if d.startswith("02"):
            f = f"02-{d[2:-4]}-{d[-4:]}"
        else:
            f = f"{d[:3]}-{d[3:-4]}-{d[-4:]}"
        mobile = d.startswith("01")
        return {"kind": "전화번호", "valid": (len(d) in (10, 11)) if mobile else True, "formatted": f,
                "detail": "휴대전화" if mobile else "일반전화", "note": ""}
    return {"kind": "", "valid": False, "formatted": body.value.strip(), "detail": "",
            "note": "사업자등록번호(10자리), 주민·법인등록번호(13자리), 카드번호, 전화번호를 확인할 수 있어요."}


# --- 10. 표 합치기 ---
TABLE_EXTS = {".xlsx", ".xlsm", ".csv", ".tsv", ".txt"}


def _read_csv(path: Path) -> list[list[str]]:
    data = path.read_bytes()
    for enc in ("utf-8-sig", "cp949"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise AppError(f"{path.name}의 글자 인코딩을 알 수 없어요.")
    delim = "\t" if path.suffix.lower() == ".tsv" or ("\t" in text.split("\n", 1)[0] and "," not in text.split("\n", 1)[0]) else ","
    return [row for row in csv.reader(io.StringIO(text), delimiter=delim)]


def read_table(path: Path, all_sheets: bool) -> list[tuple[str, list[list]]]:
    """(출처 이름, 행 목록) 목록. 엑셀은 시트마다 하나."""
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001
            raise AppError(f"{path.name}을 열 수 없어요. 암호가 걸렸거나 손상된 파일일 수 있어요.") from e
        sheets = wb.worksheets if all_sheets else wb.worksheets[:1]
        out = []
        for ws in sheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            out.append((f"{path.name} / {ws.title}" if all_sheets else path.name, rows))
        wb.close()
        return out
    return [(path.name, _read_csv(path))]


def _blank(row: list) -> bool:
    return all(v is None or str(v).strip() == "" for v in row)


class MergeIn(BaseModel):
    paths: list[str] = Field(min_length=1)
    header: bool = True  # 첫 줄이 제목 줄인지 (제목이 같은 열끼리 맞춘다)
    source_column: bool = True
    all_sheets: bool = False
    out_path: str | None = None


@router.post("/merge-tables")
def merge_tables(body: MergeIn):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    files: list[Path] = []
    for p in body.paths:
        path = Path(p)
        if path.is_dir():
            files += sorted(f for f in path.iterdir() if f.suffix.lower() in TABLE_EXTS and not f.name.startswith("~$"))
        elif path.exists():
            files.append(path)
    if not files:
        raise AppError("합칠 엑셀·CSV 파일이 없어요.")
    columns: list[str] = []
    records: list[tuple[str, dict | list]] = []
    per_file = []
    for f in files:
        for source, rows in read_table(f, body.all_sheets):
            rows = [r for r in rows if not _blank(r)]
            if not rows:
                per_file.append({"source": source, "rows": 0})
                continue
            if body.header:
                head = [str(h).strip() if h is not None and str(h).strip() else f"열{i + 1}" for i, h in enumerate(rows[0])]
                for h in head:
                    if h not in columns:
                        columns.append(h)
                for r in rows[1:]:
                    records.append((source, {h: v for h, v in zip(head, r)}))
                per_file.append({"source": source, "rows": len(rows) - 1})
            else:
                records += [(source, list(r)) for r in rows]
                per_file.append({"source": source, "rows": len(rows)})
    wb = Workbook()
    ws = wb.active
    ws.title = "합친 표"
    width = len(columns) if body.header else max((len(r) for _, r in records), default=0)
    head = (["출처"] if body.source_column else []) + (columns if body.header else [f"열{i + 1}" for i in range(width)])
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8EEF9")
    for source, rec in records:
        values = [rec.get(h) for h in columns] if isinstance(rec, dict) else list(rec) + [None] * (width - len(rec))
        ws.append(([source] if body.source_column else []) + values)
    ws.freeze_panes = "A2"
    for i, h in enumerate(head, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = min(40, max(8, len(str(h)) * 2 + 2))
    if body.out_path:
        dst = Path(body.out_path)
        dst.parent.mkdir(parents=True, exist_ok=True)
    else:
        dst = output_path(files[0].parent / f"합친표_{datetime.now():%Y%m%d_%H%M}.xlsx", "")
    wb.save(dst)
    return {"dst": str(dst), "rows": len(records), "columns": len(head), "files": per_file}
