"""회의록 내보내기: txt, md, srt, vtt, docx."""
from __future__ import annotations

from pathlib import Path

FORMATS = ("txt", "md", "srt", "vtt", "docx")


def _ts(sec: float, sep: str = ",") -> str:
    ms = int(round(sec * 1000))
    h, ms = divmod(ms, 3_600_000)
    m, ms = divmod(ms, 60_000)
    s, ms = divmod(ms, 1000)
    return f"{h:02d}:{m:02d}:{s:02d}{sep}{ms:03d}"


def _clock(sec: float) -> str:
    sec = int(sec)
    h, rem = divmod(sec, 3600)
    m, s = divmod(rem, 60)
    return f"{h:d}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"


def speaker_name(seg: dict, speakers: dict) -> str:
    sp = seg.get("speaker")
    if sp is None:
        return ""
    return speakers.get(str(sp)) or f"화자{int(sp) + 1}"


def paragraphs(segments: list[dict], speakers: dict, gap: float = 1.5) -> list[dict]:
    """화자가 바뀌거나 오래 쉬면 문단을 나눈다."""
    out: list[dict] = []
    for seg in segments:
        name = speaker_name(seg, speakers)
        if out and out[-1]["speaker"] == name and seg["start"] - out[-1]["end"] < gap:
            out[-1]["text"] += " " + seg["text"].strip()
            out[-1]["end"] = seg["end"]
        else:
            out.append({"speaker": name, "start": seg["start"], "end": seg["end"], "text": seg["text"].strip()})
    return out


def full_text(segments: list[dict]) -> str:
    return " ".join(s["text"].strip() for s in segments)


def to_txt(t: dict, timestamps: bool = True) -> str:
    lines = []
    for p in paragraphs(t["segments"], t.get("speakers", {})):
        head = []
        if timestamps:
            head.append(f"[{_clock(p['start'])}]")
        if p["speaker"]:
            head.append(f"{p['speaker']}:")
        lines.append((" ".join(head) + " " if head else "") + p["text"])
    return "\n\n".join(lines) + "\n"


def to_srt(segments: list[dict], speakers: dict | None = None) -> str:
    blocks = []
    for i, s in enumerate(segments, 1):
        name = speaker_name(s, speakers or {})
        text = f"{name}: {s['text'].strip()}" if name else s["text"].strip()
        blocks.append(f"{i}\n{_ts(s['start'])} --> {_ts(s['end'])}\n{text}\n")
    return "\n".join(blocks)


def to_vtt(segments: list[dict], speakers: dict | None = None) -> str:
    blocks = ["WEBVTT\n"]
    for s in segments:
        name = speaker_name(s, speakers or {})
        text = f"<v {name}>{s['text'].strip()}" if name else s["text"].strip()
        blocks.append(f"{_ts(s['start'], '.')} --> {_ts(s['end'], '.')}\n{text}\n")
    return "\n".join(blocks)


def to_md(t: dict, actions: list[dict] | None = None) -> str:
    out = [f"# {t.get('title') or '회의록'}", ""]
    if t.get("created_at"):
        out += [f"- 작성: {t['created_at'][:16].replace('T', ' ')}", ""]
    summary = t.get("summary") or []
    if summary:
        out += ["## 요약", *[f"- {p}" for p in summary], ""]
    if actions:
        out.append("## 할 일")
        for a in actions:
            extra = " / ".join(x for x in (a.get("assignee"), a.get("due")) if x)
            out.append(f"- [ ] {a['text']}" + (f" ({extra})" if extra else ""))
        out.append("")
    out += ["## 전체 내용", ""]
    for p in paragraphs(t["segments"], t.get("speakers", {})):
        who = f"**{p['speaker']}** " if p["speaker"] else ""
        out += [f"{who}`{_clock(p['start'])}` {p['text']}", ""]
    return "\n".join(out)


def to_docx(t: dict, path: Path, actions: list[dict] | None = None) -> None:
    import docx
    from docx.shared import Pt, RGBColor

    doc = docx.Document()
    style = doc.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.font.size = Pt(10.5)
    doc.add_heading(t.get("title") or "회의록", level=1)
    if t.get("created_at"):
        doc.add_paragraph(f"작성: {t['created_at'][:16].replace('T', ' ')}")
    summary = t.get("summary") or []
    if summary:
        doc.add_heading("요약", level=2)
        for p in summary:
            doc.add_paragraph(p, style="List Bullet")
    if actions:
        doc.add_heading("할 일", level=2)
        for a in actions:
            extra = " / ".join(x for x in (a.get("assignee"), a.get("due")) if x)
            doc.add_paragraph(f"☐ {a['text']}" + (f" ({extra})" if extra else ""))
    doc.add_heading("전체 내용", level=2)
    for p in paragraphs(t["segments"], t.get("speakers", {})):
        para = doc.add_paragraph()
        stamp = para.add_run(f"[{_clock(p['start'])}] ")
        stamp.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        if p["speaker"]:
            para.add_run(f"{p['speaker']}: ").bold = True
        para.add_run(p["text"])
    doc.save(str(path))
